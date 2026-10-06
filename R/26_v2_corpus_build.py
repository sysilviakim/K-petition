# Part 1 builds the inputs for the administrative-outcomes analysis on the
# rescored 2024 corpus (PSJ R&R, comments R1.2 / R1.6). Supersedes the corpus
# half of R/18_survey_admin_outcomes_build.py, which read the misaligned LLM
# scores in data/tidy/evaluated_data_2024_final.csv.
#
# The rescored file carries, for every 2024 public proposal, four 1-5
# dimension scores from three scorers, joined on a petition id rather than on
# file order:
#   gpt4omini_realigned   the original gpt-4o-mini scores, put back on their
#                         own petitions
#   astra                 gpt-6-astra
#   fable                 claude-fable-5-1
#
# Input:
#   evaluated_data_2024_rescored.csv   (default location: the Michigan project
#                                       folder; override with RESCORED_CSV)
#   data/screenshots/공개제안_tidy_url_appended_v2.xlsx   65 stimuli, v2 codes
# Output of part 1:
#   data/tidy/admin_outcomes_corpus2024_rescored.csv
#       one row per unique proposal: outcome variables and the twelve scores
#   data/tidy/llm_scores_65.csv
#       the 65 survey petitions: v2 author codes and the twelve scores
#
# Each proposal is also flagged for the six survey topics (keyword in the
# title or body, as in R/19_sample_benchmark_build.py). Every proposal in the
# file is already addressed to a central government agency.
#
# Deduplication and outcome coding follow script 18. Proposals are kept
# whether or not a scorer returned a score; the analysis script drops missing
# scores per scorer.
#
# Part 2 describes the short tail of the 2024 public-proposal frame from which
# the 65 survey petitions were drawn (comments R1.3 / R3.1): how common very
# short submissions are, how many of them fill the platform's required fields
# with placeholder text, and whether the administration answers them.
# Definitions follow R/19_sample_benchmark_build.py. Text statistics only; no
# LLM scores are used.
# Output:  output/v2/sensitivity/frame_short_tail.csv
#
# Usage: python3 R/26_v2_corpus_build.py   (from project root)

import csv
import os
import re
import statistics
import sys
import unicodedata
from collections import defaultdict
from datetime import date

import openpyxl

csv.field_size_limit(sys.maxsize)

RESCORED_CSV = os.environ.get(
    "RESCORED_CSV",
    os.path.expanduser(
        "~/Dropbox/Michigan/Projects/K-petition/data/tidy/"
        "evaluated_data_2024_rescored.csv"
    ),
)
STIM_XLSX = "data/screenshots/공개제안_tidy_url_appended_v2.xlsx"
OUT_CORPUS = "data/tidy/admin_outcomes_corpus2024_rescored.csv"
OUT_65 = "data/tidy/llm_scores_65.csv"

SCORERS = ["gpt4omini_realigned", "astra", "fable"]
DIMS = {
    "clarity": "clarity_specificity",
    "logic": "logic_consistency",
    "manner": "tone_manner",
    "validity": "validity_feasibility",
}
TOPICS = {
    "real_estate": "부동산",
    "low_birth_rate": "저출산",
    "pension": "연금",
    "private_education": "사교육",
    "delivery": "배달",
    "e_scooter": "킥보드",
}
SCORE_COLS = {
    f"{s}_{d}": f"{s}_{long}_score" for s in SCORERS for d, long in DIMS.items()
}


def norm(s):
    return unicodedata.normalize("NFC", (s or "").strip())


def toks(s):
    return set(re.findall(r"\S+", norm(s)))


def parse_scrape(ds):
    return date(int(ds[:4]), int(ds[4:6]), int(ds[6:8]))


def score(v):
    return "" if v in ("", "NA") else str(int(float(v)))


with open(RESCORED_CSV, encoding="utf-8-sig") as f:
    corpus = list(csv.DictReader(f))

# ---- full 2024 corpus ---------------------------------------------------------
seen = set()
outc = []
for r in corpus:
    key = (r["title"], r["date_petitioned"], r["current_issues"][:80])
    if key in seen:
        continue
    seen.add(key)
    try:
        d0 = date.fromisoformat(r["date_petitioned"])
    except ValueError:
        continue
    da = r["date_answered"]
    answered = 1 if da not in ("", "NA") else 0
    if answered:
        t = (date.fromisoformat(da) - d0).days
    else:
        t = (parse_scrape(r["date_scraped"]) - d0).days
    if t < 0:
        continue
    body = " ".join(
        [r["current_issues"], r["improvement_plan"], r["expected_effect"]]
    )
    row = {
        "pid": r["pid"],
        "status": r["status"],
        "answered": answered,
        "time": t,
        "area": r["area"],
        "nchar": len(body),
        "date_petitioned": r["date_petitioned"],
    }
    row.update({k: score(r[c]) for k, c in SCORE_COLS.items()})
    text = norm("\n".join([r["title"], body]))
    row.update({f"topic_{k}": int(kw in text) for k, kw in TOPICS.items()})
    outc.append(row)

with open(OUT_CORPUS, "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=list(outc[0].keys()))
    w.writeheader()
    w.writerows(outc)
print(f"wrote {OUT_CORPUS}: {len(outc)} unique proposals "
      f"(of {len(corpus)} rows)")
for s in SCORERS:
    n = sum(1 for o in outc if all(o[f"{s}_{d}"] != "" for d in DIMS))
    print(f"  scored on all four dimensions by {s}: {n}")
anyt = [o for o in outc if any(o[f"topic_{k}"] for k in TOPICS)]
print(f"  on a survey topic: {len(anyt)}")
for k in TOPICS:
    print(f"    {k}: {sum(o[f'topic_{k}'] for o in outc)}")

# ---- the 65 survey petitions --------------------------------------------------
wb = openpyxl.load_workbook(STIM_XLSX, read_only=True)
rows = list(wb.active.iter_rows(values_only=True))
stim = [dict(zip(rows[0], r)) for r in rows[1:]]

idx = defaultdict(list)
for r in corpus:
    idx[norm(r["title"])].append(r)

out65 = []
for s in stim:
    t, dp = norm(s["title"]), str(s["date_petitioned"])[:10]
    cands = [x for x in idx[t] if x["date_petitioned"] == dp]
    if len(cands) > 1:
        st = toks(s["text"])

        def sim(c):
            ct = toks(
                " ".join(
                    [c["current_issues"], c["improvement_plan"],
                     c["expected_effect"]]
                )
            )
            return len(st & ct) / max(len(ct), 1)

        cands.sort(key=sim, reverse=True)
    c = cands[0]
    comb = str(s["comb_fin"]).zfill(4)
    row = {"item": int(s["id"]), "pid": c["pid"]}
    row.update({f"author_{d}": int(comb[i]) for i, d in enumerate(DIMS)})
    row.update({k: score(c[col]) for k, col in SCORE_COLS.items()})
    out65.append(row)

with open(OUT_65, "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=list(out65[0].keys()))
    w.writeheader()
    w.writerows(out65)
print(f"wrote {OUT_65}: {len(out65)} petitions")


# ---- part 2: the short tail of the frame -------------------------------------

OUT_TAIL = "output/v2/sensitivity/frame_short_tail.csv"
FIELDS = ["current_issues", "improvement_plan", "expected_effect"]
ANSWERED = {"답변완료", "심사완료", "제안실현", "제안실시 (완전실시)",
            "제안실시 (일부실시)"}


def text_len(r):
    # title plus the three required fields, as in R/19_sample_benchmark_build.py
    return len(norm("\n".join([r["title"]] + [r[f] for f in FIELDS])))


def placeholder(f):
    f = norm(f)
    return len(f) <= 3 or re.fullmatch(
        r"[ㄱ-ㅎㅏ-ㅣ0-9\s.,\-~!?a-zA-Z]{0,5}", f
    ) is not None


seen, rows = set(), []
for r in corpus:
    key = (r["title"], r["date_petitioned"], r["current_issues"][:80])
    if key not in seen:
        seen.add(key)
        rows.append(r)
for r in rows:
    r["_len"] = text_len(r)
    r["_ph"] = any(placeholder(r[f]) for f in FIELDS)
    r["_ans"] = r["status"] in ANSWERED

wb = openpyxl.load_workbook(STIM_XLSX, read_only=True)
srows = list(wb.active.iter_rows(values_only=True))
slens = [len(norm(dict(zip(srows[0], r))["text"])) for r in srows[1:]]


def share(sub, key):
    return sum(1 for r in sub if r[key]) / len(sub) if sub else float("nan")


n = len(rows)
u200 = [r for r in rows if r["_len"] < 200]
u100 = [r for r in rows if r["_len"] < 100]
degen = [r for r in u100 if r["_ph"]]
degen_ids = {id(r) for r in degen}
rest = [r for r in rows if id(r) not in degen_ids]
short_ok = [r for r in u200 if not r["_ph"]]
out = [
    ("Unique proposals in the 2024 frame", n),
    ("Median text length, frame (characters)",
     statistics.median(r["_len"] for r in rows)),
    ("Median text length, 65 survey petitions (characters)",
     statistics.median(slens)),
    ("Shortest survey petition (characters)", min(slens)),
    ("Frame proposals under 200 characters", len(u200)),
    ("  share of frame", round(len(u200) / n, 4)),
    ("Frame proposals under 100 characters", len(u100)),
    ("  share of frame", round(len(u100) / n, 4)),
    ("  of those, placeholder text in a required field",
     round(share(u100, "_ph"), 4)),
    ("Degenerate submissions (under 100 with placeholder text)", len(degen)),
    ("  share answered by data collection", round(share(degen, "_ans"), 4)),
    ("All other proposals: share answered", round(share(rest, "_ans"), 4)),
    ("Short but substantive (under 200, no placeholder): share answered",
     round(share(short_ok, "_ans"), 4)),
    ("Short but substantive: number", len(short_ok)),
]
os.makedirs(os.path.dirname(OUT_TAIL), exist_ok=True)
with open(OUT_TAIL, "w", newline="") as fh:
    w = csv.writer(fh)
    w.writerow(["quantity", "value"])
    w.writerows(out)
for k, v in out:
    print(f"{k}: {v}")
